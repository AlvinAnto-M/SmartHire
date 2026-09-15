"""
Document loading utilities for extracting raw text from PDF and DOCX files.
Supports both file paths and Streamlit UploadedFile (BytesIO) objects.
"""

import io
from pathlib import Path
from typing import Union
import pypdf
import docx

def extract_text_from_pdf(file_source: Union[str, Path, io.BytesIO, bytes]) -> str:
    """
    Extracts text content from a PDF file path or byte stream.
    
    Args:
        file_source: File path (str/Path) or bytes/BytesIO buffer.
        
    Returns:
        Cleaned extracted text string.
        
    Raises:
        ValueError: If file is empty, corrupted, or unreadable.
    """
    try:
        if isinstance(file_source, (str, Path)):
            with open(file_source, "rb") as f:
                reader = pypdf.PdfReader(f)
                pages_text = [page.extract_text() or "" for page in reader.pages]
        elif isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
            reader = pypdf.PdfReader(stream)
            pages_text = [page.extract_text() or "" for page in reader.pages]
        else:
            # Handles io.BytesIO and Streamlit UploadedFile
            reader = pypdf.PdfReader(file_source)
            pages_text = [page.extract_text() or "" for page in reader.pages]

        combined_text = "\n".join(pages_text).strip()
        if not combined_text:
            raise ValueError("The uploaded PDF does not contain extractable text. It may be scanned or image-only.")
            
        return combined_text
    except Exception as e:
        if "The uploaded PDF does not contain" in str(e):
            raise
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")

def extract_text_from_docx(file_source: Union[str, Path, io.BytesIO, bytes]) -> str:
    """
    Extracts text content from a DOCX Word document.
    
    Args:
        file_source: File path (str/Path) or bytes/BytesIO buffer.
        
    Returns:
        Cleaned extracted text string.
        
    Raises:
        ValueError: If file is empty or corrupted.
    """
    try:
        if isinstance(file_source, (str, Path)):
            doc = docx.Document(file_source)
        elif isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
            doc = docx.Document(stream)
        else:
            doc = docx.Document(file_source)

        paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text if present
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    paragraphs_text.append(" | ".join(row_cells))

        combined_text = "\n".join(paragraphs_text).strip()
        if not combined_text:
            raise ValueError("The uploaded DOCX file does not contain any readable text.")
            
        return combined_text
    except Exception as e:
        if "The uploaded DOCX file does not contain" in str(e):
            raise
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

def extract_text_from_file(uploaded_file) -> str:
    """
    Dispatches to appropriate extractor based on filename extension.
    Accepts Streamlit UploadedFile, file path, or named buffer.
    """
    filename = getattr(uploaded_file, "name", str(uploaded_file)).lower()
    
    if filename.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    elif filename.endswith(".txt"):
        if hasattr(uploaded_file, "read"):
            content = uploaded_file.read()
            if isinstance(content, bytes):
                return content.decode("utf-8", errors="ignore").strip()
            return str(content).strip()
        with open(uploaded_file, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
    else:
        raise ValueError("Unsupported file format. Please upload a PDF (.pdf), Word document (.docx), or plain text (.txt).")
